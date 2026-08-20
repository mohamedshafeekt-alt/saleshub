"""Bulk lead import: build a downloadable sample template (xlsx/csv) and
parse an uploaded file of either format into created Lead rows.

Every imported lead is owned by whoever uploads the file -- there is no
owner column in the sheet, and no way to assign a row to someone else via
the file's contents (an "owner_email"-like column, if present in an
uploaded file, is not one of IMPORT_COLUMNS and is simply dropped by the
per-row normalization in _parse_rows).

Per-row errors (missing required field, invalid enum value, duplicate email)
are collected instead of failing the whole batch on one bad row. Each
successful row is committed immediately rather than once at the end:
create_lead's duplicate-email handling calls db.rollback() on IntegrityError,
and that rolls back the whole shared session -- not just the failing row's
own uncommitted work -- so leaving earlier rows uncommitted would let one bad
row silently wipe out every lead already created earlier in the same batch.
Committing per row means there is never anything earlier left to lose.
"""

import csv
import io
from typing import Literal

from fastapi import BackgroundTasks
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.lead import Lead
from app.models.user import User
from app.schemas.lead import LeadUpsert
from app.schemas.lead_import import LeadImportResult, LeadImportRowError
from app.services.email.sender import EmailSender
from app.services.lead_service import DuplicateLeadEmailError, create_lead

IMPORT_COLUMNS = [
    "first_name",
    "last_name",
    "company",
    "domain",
    "job_title",
    "linkedin_url",
    "email",
    "phone",
    "source",
    "status",
    "follow_up_note",
]

_EXAMPLE_ROW = {
    "first_name": "vishnu",
    "last_name": "ram",
    "company": "IB",
    "domain": "ib.com",
    "job_title": "VP of Sales",
    "linkedin_url": "https://linkedin.com/in/vishnuram",
    "email": "vishnu.ram@ib.com",
    "phone": "1234567890",
    "source": "website",
    "status": "not_contacted",
    "follow_up_note": "follow up note",
}


class LeadImportFileError(Exception):
    """Raised when the uploaded file can't be parsed as the expected format
    (wrong/corrupt content despite a matching extension, bad text encoding,
    no header row, etc)."""


def build_lead_import_template(fmt: Literal["xlsx", "csv"]) -> bytes:
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=IMPORT_COLUMNS)
        writer.writeheader()
        writer.writerow(_EXAMPLE_ROW)
        return buffer.getvalue().encode("utf-8")

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # a freshly created Workbook always has an active sheet
    sheet.title = "Leads"
    sheet.append(IMPORT_COLUMNS)
    sheet.append([_EXAMPLE_ROW[column] for column in IMPORT_COLUMNS])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _parse_csv_rows(file_bytes: bytes) -> list[dict[str, str | None]]:
    text = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [{(key or "").strip().lower(): _cell_str(value) for key, value in row.items()} for row in reader]


def _parse_xlsx_rows(file_bytes: bytes) -> list[dict[str, str | None]]:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    assert sheet is not None  # a loaded workbook always has an active sheet
    rows_iter = sheet.iter_rows(values_only=True)
    header = [(_cell_str(cell) or "").lower() for cell in next(rows_iter)]
    rows = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        rows.append({header[i]: _cell_str(raw_row[i]) for i in range(len(header))})
    return rows


def _parse_rows(file_bytes: bytes, filename: str) -> list[dict[str, str | None]]:
    try:
        rows = _parse_csv_rows(file_bytes) if filename.lower().endswith(".csv") else _parse_xlsx_rows(file_bytes)
        # Normalize every row to exactly IMPORT_COLUMNS: a column absent from
        # the uploaded file's header (renamed, reordered away, or dropped)
        # becomes None here instead of a KeyError wherever that field is read,
        # and any column NOT in IMPORT_COLUMNS (e.g. a leftover owner_email
        # from an old template) is dropped rather than acted on.
        return [{column: row.get(column) for column in IMPORT_COLUMNS} for row in rows]
    except LeadImportFileError:
        raise
    except Exception as exc:
        raise LeadImportFileError(f"Could not read uploaded file: {exc}") from exc


async def import_leads(
    db: AsyncSession,
    file_bytes: bytes,
    filename: str,
    requester: User,
    email_sender: EmailSender,
    background_tasks: BackgroundTasks | None = None,
) -> LeadImportResult:
    """requester is the user uploading the file: every created lead is owned
    by them -- see module docstring.

    background_tasks is threaded straight through to each row's create_lead
    call so notify-on-create emails defer to it instead of being awaited
    inline -- without it, create_lead sends one SMTP round-trip per
    notify-on-create admin PER ROW, which is what makes a multi-row import
    slow (same failure mode already fixed for single-lead create)."""
    rows = _parse_rows(file_bytes, filename)

    row_emails = {row["email"].lower() for row in rows if row["email"]}
    existing_emails: set[str] = set()
    if row_emails:
        result = await db.execute(select(func.lower(Lead.email)).where(func.lower(Lead.email).in_(row_emails)))
        existing_emails = set(result.scalars().all())

    seen_in_file: set[str] = set()
    created = 0
    errors: list[LeadImportRowError] = []

    for index, row in enumerate(rows, start=2):  # header is row 1
        email = row.get("email")
        email_key = email.lower() if email else None

        if email_key and email_key in existing_emails:
            errors.append(LeadImportRowError(row=index, error=f"Email already exists: {email}"))
            continue
        if email_key and email_key in seen_in_file:
            errors.append(LeadImportRowError(row=index, error=f"Duplicate email in file: {email}"))
            continue

        try:
            data = LeadUpsert(
                first_name=row.get("first_name"),
                last_name=row.get("last_name"),
                company=row.get("company"),
                domain=row.get("domain"),
                job_title=row.get("job_title"),
                linkedin_url=row.get("linkedin_url"),
                email=row.get("email"),
                phone=row.get("phone"),
                source=row.get("source"),
                status=row.get("status"),
                owner_id=requester.id,
                follow_up_note=row.get("follow_up_note"),
            )
        except ValidationError as exc:
            errors.append(LeadImportRowError(row=index, error=str(exc)))
            continue

        try:
            await create_lead(db, data, email_sender, requester=requester, background_tasks=background_tasks)
        except (DuplicateLeadEmailError, IntegrityError) as exc:
            errors.append(LeadImportRowError(row=index, error=str(exc)))
            continue

        # Committed immediately (rather than once at the end): create_lead's
        # own duplicate-email rollback undoes the whole session, not just the
        # failing row, so a later row's failure must never find this row's
        # work still sitting uncommitted -- see module docstring.
        await db.commit()

        if email_key:
            seen_in_file.add(email_key)
        created += 1

    return LeadImportResult(created=created, errors=errors)
