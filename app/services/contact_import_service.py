"""Bulk contact import: build a downloadable sample template (xlsx/csv) and
parse an uploaded file of either format into created standalone Contact rows.

Imported contacts are never linked to an Account -- there is no
company/account column in the sheet. Link an imported contact to an Account
afterward via the existing Account Detail page's Add Contact flow
(app.services.contact_account_service).

Per-row errors (missing required field, invalid value, duplicate email) are
collected instead of failing the whole batch on one bad row. Each successful
row is committed immediately rather than once at the end: create_contact's
duplicate-email handling calls db.rollback() on IntegrityError, and that
rolls back the whole shared session -- not just the failing row's own
uncommitted work -- so leaving earlier rows uncommitted would let one bad row
silently wipe out every contact already created earlier in the same batch.
Committing per row means there is never anything earlier left to lose.
"""

import csv
import io
from typing import Literal

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contact import Contact
from app.schemas.contact import ContactCreate
from app.schemas.contact_import import ContactImportResult, ContactImportRowError
from app.services.contact_service import DuplicateContactEmailError, create_contact

IMPORT_COLUMNS = [
    "first_name",
    "last_name",
    "email",
    "phone",
    "alternate_phone",
    "job_title",
    "linkedin_url",
]

_EXAMPLE_ROW = {
    "first_name": "vishnu",
    "last_name": "ram",
    "email": "vishnu.ram@ib.com",
    "phone": "1234567890",
    "alternate_phone": "9876543210",
    "job_title": "VP of Sales",
    "linkedin_url": "https://linkedin.com/in/vishnuram",
}


class ContactImportFileError(Exception):
    """Raised when the uploaded file can't be parsed as the expected format
    (wrong/corrupt content despite a matching extension, bad text encoding,
    no header row, etc)."""


def build_contact_import_template(fmt: Literal["xlsx", "csv"]) -> bytes:
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=IMPORT_COLUMNS)
        writer.writeheader()
        writer.writerow(_EXAMPLE_ROW)
        return buffer.getvalue().encode("utf-8")

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # a freshly created Workbook always has an active sheet
    sheet.title = "Contacts"
    sheet.append(IMPORT_COLUMNS)
    sheet.append([_EXAMPLE_ROW[column] for column in IMPORT_COLUMNS])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _parse_csv_rows(file_bytes: bytes) -> list[dict[str, str | None]]:
    text = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [{(key or "").strip().lower(): _cell_str(value) for key, value in row.items()} for row in reader]


def _parse_xlsx_rows(file_bytes: bytes) -> list[dict[str, str | None]]:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    assert sheet is not None  # a loaded workbook always has an active sheet
    rows_iter = sheet.iter_rows(values_only=True)
    header = [(_cell_str(cell) or "").lower() for cell in next(rows_iter)]
    rows = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        rows.append({header[i]: _cell_str(raw_row[i]) for i in range(len(header))})
    return rows


def _parse_rows(file_bytes: bytes, filename: str) -> list[dict[str, str | None]]:
    try:
        rows = _parse_csv_rows(file_bytes) if filename.lower().endswith(".csv") else _parse_xlsx_rows(file_bytes)
        # Normalize every row to exactly IMPORT_COLUMNS: a column absent from
        # the uploaded file's header (renamed, reordered away, or dropped)
        # becomes None here instead of a KeyError wherever that field is read,
        # and any column NOT in IMPORT_COLUMNS is dropped rather than acted on.
        return [{column: row.get(column) for column in IMPORT_COLUMNS} for row in rows]
    except ContactImportFileError:
        raise
    except Exception as exc:
        raise ContactImportFileError(f"Could not read uploaded file: {exc}") from exc


async def import_contacts(db: AsyncSession, file_bytes: bytes, filename: str) -> ContactImportResult:
    rows = _parse_rows(file_bytes, filename)

    row_emails = {row["email"].lower() for row in rows if row["email"]}
    existing_emails: set[str] = set()
    if row_emails:
        result = await db.execute(
            select(func.lower(Contact.email)).where(func.lower(Contact.email).in_(row_emails))
        )
        existing_emails = set(result.scalars().all())

    seen_in_file: set[str] = set()
    created = 0
    errors: list[ContactImportRowError] = []

    for index, row in enumerate(rows, start=2):  # header is row 1
        email = row.get("email")
        email_key = email.lower() if email else None

        if email_key and email_key in existing_emails:
            errors.append(ContactImportRowError(row=index, error=f"Email already exists: {email}"))
            continue
        if email_key and email_key in seen_in_file:
            errors.append(ContactImportRowError(row=index, error=f"Duplicate email in file: {email}"))
            continue

        try:
            data = ContactCreate(
                first_name=row.get("first_name"),
                last_name=row.get("last_name"),
                email=row.get("email"),
                phone=row.get("phone"),
                alternate_phone=row.get("alternate_phone"),
                job_title=row.get("job_title"),
                linkedin_url=row.get("linkedin_url"),
            )
        except ValidationError as exc:
            errors.append(ContactImportRowError(row=index, error=str(exc)))
            continue

        try:
            await create_contact(db, data)
        except (DuplicateContactEmailError, IntegrityError) as exc:
            errors.append(ContactImportRowError(row=index, error=str(exc)))
            continue

        # Committed immediately (rather than once at the end): create_contact's
        # own duplicate-email rollback undoes the whole session, not just the
        # failing row, so a later row's failure must never find this row's
        # work still sitting uncommitted -- see module docstring.
        await db.commit()

        if email_key:
            seen_in_file.add(email_key)
        created += 1

    return ContactImportResult(created=created, errors=errors)
