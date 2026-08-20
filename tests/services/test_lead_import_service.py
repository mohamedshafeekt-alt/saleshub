"""app.services.lead_import_service: template generation for both formats,
and per-row import of an uploaded xlsx/csv into Lead + primary LeadContact
rows, collecting row-level errors instead of failing the whole batch.

Covers: template round-trips for both formats; valid row create (xlsx + csv);
missing required field; invalid source enum; duplicate email against an
existing DB lead; duplicate email within the same file; every created lead is
owned by the uploader (an "owner_email"-like column in the file, if present,
is ignored rather than treated as an owner override); a file missing the
"email" header column; a corrupt/malformed upload; and that a later row's
duplicate-email failure doesn't wipe out an earlier row already created in
the same batch.
"""

import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.lead import Lead
from app.models.user import User
from app.services.lead_import_service import (
    IMPORT_COLUMNS,
    LeadImportFileError,
    build_lead_import_template,
    import_leads,
)
from tests.support.roles import UserRole, role_id_for


class FakeEmailSender:
    def __init__(self) -> None:
        self.calls: list[dict] = []

    async def send(self, to: str, subject: str, body: str) -> None:
        self.calls.append({"to": to, "subject": subject, "body": body})


async def _make_user(db_session: AsyncSession, email: str, role: UserRole = UserRole.SALES_REP) -> User:
    user = User(email=email, hashed_password="x", first_name="Test", role_id=await role_id_for(db_session, role))
    db_session.add(user)
    await db_session.flush()
    return user


def _xlsx_bytes(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(IMPORT_COLUMNS)
    for row in rows:
        sheet.append([row.get(col, "") for col in IMPORT_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _csv_bytes(rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=IMPORT_COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow({col: row.get(col, "") for col in IMPORT_COLUMNS})
    return buffer.getvalue().encode("utf-8")


async def test_import_leads_defers_notify_emails_to_background_task(db_session: AsyncSession):
    """Every row's create_lead call must defer notify-on-create emails to a
    background_tasks handle, same as the single-lead create route -- an
    import doesn't pass background_tasks through, each row awaits SMTP
    sends inline for every notify-on-create admin, which is what makes a
    multi-row import slow (N rows x M admins sequential SMTP round-trips
    inside the request)."""
    from fastapi import BackgroundTasks

    uploader = await _make_user(db_session, "uploader-bg@acme.com")
    admin = await _make_user(db_session, "bg-import-admin@example.com", UserRole.ADMIN)
    fake_sender = FakeEmailSender()
    background_tasks = BackgroundTasks()
    row = {"first_name": "Jane", "company": "Acme Corp", "email": "bg-import-lead@acme.com", "source": "website"}

    result = await import_leads(
        db_session, _csv_bytes([row]), "leads.csv", uploader, fake_sender, background_tasks=background_tasks
    )

    assert result.created == 1
    assert fake_sender.calls == []  # not sent yet -- deferred, not blocking the import request
    await background_tasks()
    assert {call["to"] for call in fake_sender.calls} == {admin.email}


def test_build_lead_import_template_xlsx_has_header_and_example_row():
    content = build_lead_import_template("xlsx")

    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    assert header == IMPORT_COLUMNS
    assert sheet.max_row == 2


def test_build_lead_import_template_csv_has_header_and_example_row():
    content = build_lead_import_template("csv")

    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))

    assert rows[0] == IMPORT_COLUMNS
    assert len(rows) == 2


async def test_import_leads_creates_valid_row_from_xlsx(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-xlsx@acme.com")
    row = {"first_name": "Jane", "company": "Acme Corp", "email": "jane.import.xlsx@acme.com", "source": "website"}

    result = await import_leads(db_session, _xlsx_bytes([row]), "leads.xlsx", uploader, FakeEmailSender())

    assert result.created == 1
    assert result.errors == []
    created = (
        await db_session.execute(select(Lead).where(Lead.email == "jane.import.xlsx@acme.com"))
    ).scalar_one()
    assert created.company == "Acme Corp"


async def test_import_leads_creates_valid_row_from_csv(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-csv@acme.com")
    row = {"first_name": "Jane", "company": "Acme Corp", "email": "jane.import.csv@acme.com", "source": "website"}

    result = await import_leads(db_session, _csv_bytes([row]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 1
    assert result.errors == []
    created = (
        await db_session.execute(select(Lead).where(Lead.email == "jane.import.csv@acme.com"))
    ).scalar_one()
    assert created.company == "Acme Corp"


async def test_import_leads_reports_missing_required_field(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-missing-field@acme.com")
    row = {"first_name": "Jane"}  # missing company/email/source

    result = await import_leads(db_session, _csv_bytes([row]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 0
    assert len(result.errors) == 1
    assert result.errors[0].row == 2


async def test_import_leads_reports_invalid_source(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-bad-source@acme.com")
    row = {
        "first_name": "Jane",
        "company": "Acme",
        "email": "jane.badsource@acme.com",
        "source": "not_a_real_source",
    }

    result = await import_leads(db_session, _csv_bytes([row]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 0
    assert len(result.errors) == 1


async def test_import_leads_reports_duplicate_against_existing_lead(db_session: AsyncSession, make_lead):
    uploader = await _make_user(db_session, "uploader-dup-existing@acme.com")
    await make_lead(email="existing-import@acme.com")
    row = {"first_name": "Jane", "company": "Acme", "email": "existing-import@acme.com", "source": "website"}

    result = await import_leads(db_session, _csv_bytes([row]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 0
    assert len(result.errors) == 1
    assert "existing-import@acme.com" in result.errors[0].error


async def test_import_leads_reports_duplicate_within_file(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-dup-file@acme.com")
    row = {"first_name": "Jane", "company": "Acme", "email": "dup-import@acme.com", "source": "website"}

    result = await import_leads(db_session, _csv_bytes([row, dict(row)]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 1
    assert len(result.errors) == 1
    assert result.errors[0].row == 3


async def test_import_leads_owner_is_always_the_uploader(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-default-owner@acme.com")
    row = {"first_name": "Jane", "company": "Acme", "email": "owner-is-uploader@acme.com", "source": "website"}

    result = await import_leads(db_session, _csv_bytes([row]), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 1
    created = (
        await db_session.execute(select(Lead).where(Lead.email == "owner-is-uploader@acme.com"))
    ).scalar_one()
    assert created.owner_id == uploader.id


async def test_import_leads_ignores_an_owner_email_column_if_present(db_session: AsyncSession):
    """The sheet no longer has an owner_email column at all, but if someone
    reuses an old template (or adds the column back by hand) it must not be
    treated as an owner override -- the uploader stays the owner regardless."""
    uploader = await _make_user(db_session, "uploader-ignores-owner-email@acme.com")
    other_user = await _make_user(db_session, "someone-else@acme.com")
    buffer = io.StringIO()
    fieldnames = [*IMPORT_COLUMNS, "owner_email"]
    writer = csv.DictWriter(buffer, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerow(
        {
            "first_name": "Jane",
            "company": "Acme",
            "email": "ignores-owner-email-col@acme.com",
            "source": "website",
            "owner_email": other_user.email,
        }
    )
    csv_bytes = buffer.getvalue().encode("utf-8")

    result = await import_leads(db_session, csv_bytes, "leads.csv", uploader, FakeEmailSender())

    assert result.created == 1
    created = (
        await db_session.execute(select(Lead).where(Lead.email == "ignores-owner-email-col@acme.com"))
    ).scalar_one()
    assert created.owner_id == uploader.id


async def test_import_leads_missing_email_header_column_reports_row_error(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-missing-header@acme.com")
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["first_name", "company", "source"])
    writer.writeheader()
    writer.writerow({"first_name": "Jane", "company": "Acme", "source": "website"})
    csv_bytes = buffer.getvalue().encode("utf-8")

    result = await import_leads(db_session, csv_bytes, "leads.csv", uploader, FakeEmailSender())

    assert result.created == 0
    assert len(result.errors) == 1
    assert result.errors[0].row == 2


async def test_import_leads_raises_for_corrupt_xlsx_upload(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-corrupt-xlsx@acme.com")

    with pytest.raises(LeadImportFileError):
        await import_leads(db_session, b"not a real xlsx file", "leads.xlsx", uploader, FakeEmailSender())


async def test_import_leads_raises_for_non_utf8_csv_upload(db_session: AsyncSession):
    uploader = await _make_user(db_session, "uploader-bad-encoding@acme.com")
    csv_bytes = "first_name,company,email,source\nJosé,Acme,jose@acme.com,website\n".encode("latin-1")

    with pytest.raises(LeadImportFileError):
        await import_leads(db_session, csv_bytes, "leads.csv", uploader, FakeEmailSender())


async def test_import_leads_earlier_row_survives_later_duplicate_email_failure(db_session: AsyncSession):
    """create_lead's duplicate-email handling rolls back the whole shared
    session on IntegrityError, not just the failing row -- verified directly
    against the test DB (see the plan/analysis notes). Committing each
    successful row immediately (rather than once at the end) is what keeps
    an earlier row from being silently wiped out by a later row's failure."""
    uploader = await _make_user(db_session, "uploader-survives-dup@acme.com")
    rows = [
        {"first_name": "First", "company": "Acme", "email": "first-row@acme.com", "source": "website"},
        {"first_name": "Dup", "company": "Acme", "email": "first-row@acme.com", "source": "website"},
        {"first_name": "Third", "company": "Acme", "email": "third-row@acme.com", "source": "website"},
    ]

    result = await import_leads(db_session, _csv_bytes(rows), "leads.csv", uploader, FakeEmailSender())

    assert result.created == 2
    assert len(result.errors) == 1
    emails = (
        await db_session.execute(
            select(Lead.email).where(Lead.email.in_(["first-row@acme.com", "third-row@acme.com"]))
        )
    ).scalars().all()
    assert sorted(emails) == ["first-row@acme.com", "third-row@acme.com"]
