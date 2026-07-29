"""app.services.contact_import_service: template generation for both formats,
and per-row import of an uploaded xlsx/csv into standalone Contact rows,
collecting row-level errors instead of failing the whole batch.

Covers: template round-trips for both formats; valid row create (xlsx + csv);
missing required field; duplicate email against an existing DB contact;
duplicate email within the same file; a file missing the "email" header
column; a corrupt/malformed upload; and that a later row's duplicate-email
failure doesn't wipe out an earlier row already created in the same batch.
"""

import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contact import Contact
from app.schemas.contact import ContactCreate
from app.services.contact_import_service import (
    IMPORT_COLUMNS,
    ContactImportFileError,
    build_contact_import_template,
    import_contacts,
)
from app.services.contact_service import create_contact


def _xlsx_bytes(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(IMPORT_COLUMNS)
    for row in rows:
        sheet.append([row.get(col, "") for col in IMPORT_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _csv_bytes(rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=IMPORT_COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow({col: row.get(col, "") for col in IMPORT_COLUMNS})
    return buffer.getvalue().encode("utf-8")


def test_build_contact_import_template_xlsx_has_header_and_example_row():
    content = build_contact_import_template("xlsx")

    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    assert header == IMPORT_COLUMNS
    assert sheet.max_row == 2


def test_build_contact_import_template_csv_has_header_and_example_row():
    content = build_contact_import_template("csv")

    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))

    assert rows[0] == IMPORT_COLUMNS
    assert len(rows) == 2


async def test_import_contacts_creates_valid_row_from_xlsx(db_session: AsyncSession):
    row = {"first_name": "Jane", "last_name": "Doe", "email": "jane.import.xlsx@acme.com", "job_title": "VP Sales"}

    result = await import_contacts(db_session, _xlsx_bytes([row]), "contacts.xlsx")

    assert result.created == 1
    assert result.errors == []
    created = (
        await db_session.execute(select(Contact).where(Contact.email == "jane.import.xlsx@acme.com"))
    ).scalar_one()
    assert created.job_title == "VP Sales"


async def test_import_contacts_creates_valid_row_from_csv(db_session: AsyncSession):
    row = {"first_name": "Jane", "last_name": "Doe", "email": "jane.import.csv@acme.com", "job_title": "VP Sales"}

    result = await import_contacts(db_session, _csv_bytes([row]), "contacts.csv")

    assert result.created == 1
    assert result.errors == []
    created = (
        await db_session.execute(select(Contact).where(Contact.email == "jane.import.csv@acme.com"))
    ).scalar_one()
    assert created.job_title == "VP Sales"


async def test_import_contacts_reports_missing_required_field(db_session: AsyncSession):
    row = {"first_name": "Jane"}  # missing email

    result = await import_contacts(db_session, _csv_bytes([row]), "contacts.csv")

    assert result.created == 0
    assert len(result.errors) == 1
    assert result.errors[0].row == 2


async def test_import_contacts_reports_invalid_linkedin_url(db_session: AsyncSession):
    row = {"first_name": "Jane", "email": "jane.badlinkedin@acme.com", "linkedin_url": "linkedin.com/in/jane"}

    result = await import_contacts(db_session, _csv_bytes([row]), "contacts.csv")

    assert result.created == 0
    assert len(result.errors) == 1


async def test_import_contacts_reports_duplicate_against_existing_contact(db_session: AsyncSession):
    await create_contact(
        db_session, ContactCreate(first_name="Existing", email="existing-import@acme.com")
    )
    row = {"first_name": "Jane", "email": "existing-import@acme.com"}

    result = await import_contacts(db_session, _csv_bytes([row]), "contacts.csv")

    assert result.created == 0
    assert len(result.errors) == 1
    assert "existing-import@acme.com" in result.errors[0].error


async def test_import_contacts_reports_duplicate_within_file(db_session: AsyncSession):
    row = {"first_name": "Jane", "email": "dup-import@acme.com"}

    result = await import_contacts(db_session, _csv_bytes([row, dict(row)]), "contacts.csv")

    assert result.created == 1
    assert len(result.errors) == 1
    assert result.errors[0].row == 3


async def test_import_contacts_missing_email_header_column_reports_row_error(db_session: AsyncSession):
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["first_name", "last_name"])
    writer.writeheader()
    writer.writerow({"first_name": "Jane", "last_name": "Doe"})
    csv_bytes = buffer.getvalue().encode("utf-8")

    result = await import_contacts(db_session, csv_bytes, "contacts.csv")

    assert result.created == 0
    assert len(result.errors) == 1
    assert result.errors[0].row == 2


async def test_import_contacts_raises_for_corrupt_xlsx_upload(db_session: AsyncSession):
    with pytest.raises(ContactImportFileError):
        await import_contacts(db_session, b"not a real xlsx file", "contacts.xlsx")


async def test_import_contacts_raises_for_non_utf8_csv_upload(db_session: AsyncSession):
    csv_bytes = "first_name,email\nJosé,jose@acme.com\n".encode("latin-1")

    with pytest.raises(ContactImportFileError):
        await import_contacts(db_session, csv_bytes, "contacts.csv")


async def test_import_contacts_earlier_row_survives_later_duplicate_email_failure(db_session: AsyncSession):
    """create_contact's duplicate-email handling rolls back the whole shared
    session on IntegrityError, not just the failing row. Committing each
    successful row immediately (rather than once at the end) is what keeps
    an earlier row from being silently wiped out by a later row's failure."""
    rows = [
        {"first_name": "First", "email": "first-row@acme.com"},
        {"first_name": "Dup", "email": "first-row@acme.com"},
        {"first_name": "Third", "email": "third-row@acme.com"},
    ]

    result = await import_contacts(db_session, _csv_bytes(rows), "contacts.csv")

    assert result.created == 2
    assert len(result.errors) == 1
    emails = (
        await db_session.execute(
            select(Contact.email).where(Contact.email.in_(["first-row@acme.com", "third-row@acme.com"]))
        )
    ).scalars().all()
    assert set(emails) == {"first-row@acme.com", "third-row@acme.com"}
