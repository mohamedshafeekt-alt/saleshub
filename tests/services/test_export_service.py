"""app.services.export_service: rows_to_xlsx builds an in-memory workbook
from a header list + row-of-values list, with no per-resource logic."""

import io

import openpyxl

from app.services.export_service import field_value_sheet, rows_to_xlsx, sheets_to_xlsx


def test_rows_to_xlsx_writes_header_and_rows() -> None:
    headers = ["Name", "Value"]
    rows = [["Alpha", 1], ["Beta", 2]]

    buffer = rows_to_xlsx(headers, rows, sheet_name="Widgets")

    assert isinstance(buffer, io.BytesIO)
    assert buffer.tell() == 0

    workbook = openpyxl.load_workbook(buffer)
    sheet = workbook.active
    assert sheet.title == "Widgets"
    assert [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))] == headers
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert data_rows == [("Alpha", 1), ("Beta", 2)]


def test_rows_to_xlsx_with_no_rows_still_writes_header() -> None:
    buffer = rows_to_xlsx(["A", "B"], [])

    workbook = openpyxl.load_workbook(buffer)
    sheet = workbook.active
    assert sheet.title == "Sheet1"
    assert sheet.max_row == 1


def test_sheets_to_xlsx_writes_multiple_named_sheets() -> None:
    buffer = sheets_to_xlsx(
        [
            ("First", ["A", "B"], [[1, 2]]),
            ("Second", ["C"], [["x"], ["y"]]),
        ]
    )

    workbook = openpyxl.load_workbook(buffer)
    assert workbook.sheetnames == ["First", "Second"]
    first = workbook["First"]
    assert [cell.value for cell in next(first.iter_rows(min_row=1, max_row=1))] == ["A", "B"]
    assert list(first.iter_rows(min_row=2, values_only=True)) == [(1, 2)]
    second = workbook["Second"]
    assert list(second.iter_rows(min_row=2, values_only=True)) == [("x",), ("y",)]


def test_rows_to_xlsx_is_a_single_sheet_wrapper_around_sheets_to_xlsx() -> None:
    buffer = rows_to_xlsx(["Name"], [["Alpha"]], sheet_name="Only")

    workbook = openpyxl.load_workbook(buffer)
    assert workbook.sheetnames == ["Only"]


def test_field_value_sheet_builds_a_two_column_sheet_tuple() -> None:
    sheet_name, headers, rows = field_value_sheet("Lead", {"Company": "Acme", "Owner": None})

    assert sheet_name == "Lead"
    assert headers == ["Field", "Value"]
    assert rows == [["Company", "Acme"], ["Owner", None]]
