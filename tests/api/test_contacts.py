"""HTTP-level contract for /api/v1/contacts.

Standalone Contact CRUD -- role-gated only (401/403 via the router's
require_role dependency), no ownership scoping (Contact has no owner_id and
no single owning account -- see contact_service.py's module docstring).
Covers: 201 create as each allowed role, 422 missing required field, 401 no
auth, 403 for Delivery SME, 200 successful get/PATCH partial update, 204
successful DELETE, 404 for a nonexistent contact id.

Also covers the Contacts List (GET /contacts, filters + pagination), Contact
Overview (GET /contacts/{id}/overview), and Contact Deals (GET
/contacts/{id}/deals) screens.
"""

import io
import uuid

import openpyxl
from httpx import AsyncClient

from tests.support.roles import UserRole

CONTACTS_URL = "/api/v1/contacts"


def _contact_payload(**overrides) -> dict:
    # email defaults to a random unique address (Contact.email is NOT NULL
    # and globally unique) rather than a fixed constant, so calls in
    # different tests don't collide unless they explicitly pass the same
    # email on purpose (e.g. the duplicate-email tests).
    payload = {"first_name": "Jane", "email": f"jane-{uuid.uuid4().hex[:12]}@example.com"}
    payload.update(overrides)
    return payload


async def test_create_contact_as_sales_rep_returns_201(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-create-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        CONTACTS_URL, json=_contact_payload(first_name="Jane", last_name="Doe"), headers=headers
    )

    assert response.status_code == 201
    body = response.json()
    assert body["first_name"] == "Jane"
    assert body["last_name"] == "Doe"
    assert "id" in body


async def test_create_contact_as_sales_manager_returns_201(
    client: AsyncClient, make_user, auth_headers
):
    manager = await make_user(email="manager-create-contact@example.com", role=UserRole.SALES_MANAGER)
    headers = auth_headers(manager)

    response = await client.post(CONTACTS_URL, json=_contact_payload(), headers=headers)

    assert response.status_code == 201


async def test_create_contact_as_admin_returns_201(client: AsyncClient, make_user, auth_headers):
    admin = await make_user(email="admin-create-contact@example.com", role=UserRole.ADMIN)
    headers = auth_headers(admin)

    response = await client.post(CONTACTS_URL, json=_contact_payload(), headers=headers)

    assert response.status_code == 201


async def test_create_contact_with_linkedin_and_alternate_phone_returns_201(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-create-contact-extra@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        CONTACTS_URL,
        json=_contact_payload(linkedin_url="https://linkedin.com/in/jane", alternate_phone="555-0002"),
        headers=headers,
    )

    assert response.status_code == 201
    body = response.json()
    assert body["linkedin_url"] == "https://linkedin.com/in/jane"
    assert body["alternate_phone"] == "555-0002"


async def test_create_contact_returns_409_when_email_already_exists(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-create-contact-dup-email@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    first = await client.post(
        CONTACTS_URL,
        json=_contact_payload(first_name="Original", email="dup-contact-api@example.com"),
        headers=headers,
    )
    assert first.status_code == 201

    response = await client.post(
        CONTACTS_URL,
        json=_contact_payload(first_name="Different Name", email="dup-contact-api@example.com"),
        headers=headers,
    )

    assert response.status_code == 409


async def test_create_contact_missing_first_name_returns_422(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-missing-name-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    payload = _contact_payload()
    del payload["first_name"]

    response = await client.post(CONTACTS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_contact_missing_email_returns_422(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-missing-email-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    payload = _contact_payload()
    del payload["email"]

    response = await client.post(CONTACTS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_contact_no_auth_header_returns_401(client: AsyncClient):
    response = await client.post(CONTACTS_URL, json=_contact_payload())

    assert response.status_code == 401


async def test_create_contact_as_delivery_sme_returns_403(
    client: AsyncClient, make_user, auth_headers
):
    sme = await make_user(email="sme-create-contact@example.com", role=UserRole.DELIVERY_SME)
    headers = auth_headers(sme)

    response = await client.post(CONTACTS_URL, json=_contact_payload(), headers=headers)

    assert response.status_code == 403


async def test_get_contact_returns_200(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-get-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    create_response = await client.post(
        CONTACTS_URL, json=_contact_payload(first_name="Getable"), headers=headers
    )
    contact_id = create_response.json()["id"]

    response = await client.get(f"{CONTACTS_URL}/{contact_id}", headers=headers)

    assert response.status_code == 200
    assert response.json()["id"] == contact_id


async def test_get_contact_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-get-404-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/999999", headers=headers)

    assert response.status_code == 404


async def test_update_contact_partial_patch_returns_200(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-patch-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    create_response = await client.post(
        CONTACTS_URL, json=_contact_payload(first_name="Old", job_title="Old Title"), headers=headers
    )
    contact_id = create_response.json()["id"]

    response = await client.patch(f"{CONTACTS_URL}/{contact_id}", json={"first_name": "New"}, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert body["first_name"] == "New"
    assert body["job_title"] == "Old Title"


async def test_update_contact_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-patch-404-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.patch(f"{CONTACTS_URL}/999999", json={"first_name": "New"}, headers=headers)

    assert response.status_code == 404


async def test_delete_contact_returns_204(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-delete-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    create_response = await client.post(CONTACTS_URL, json=_contact_payload(), headers=headers)
    contact_id = create_response.json()["id"]

    response = await client.delete(f"{CONTACTS_URL}/{contact_id}", headers=headers)

    assert response.status_code == 204

    follow_up = await client.get(f"{CONTACTS_URL}/{contact_id}", headers=headers)
    assert follow_up.status_code == 404


async def test_delete_contact_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-delete-404-contact@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.delete(f"{CONTACTS_URL}/999999", headers=headers)

    assert response.status_code == 404


# --- GET /contacts (list) -----------------------------------------------------


async def test_list_contacts_route_returns_200_paginated(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact
):
    rep = await make_user(email="rep-list-contacts-api@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    account = await make_account(owner_id=rep.id, company="List API Co")
    contact = await make_contact(account_id=account.id, first_name="Listed")

    response = await client.get(CONTACTS_URL, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert body["total"] >= 1
    assert any(item["id"] == contact.id for item in body["items"])


async def test_list_contacts_route_filters_by_owner_id(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact
):
    owner_a = await make_user(email="owner-a-list-contacts-api@example.com", role=UserRole.SALES_REP)
    owner_b = await make_user(email="owner-b-list-contacts-api@example.com", role=UserRole.SALES_REP)
    account_a = await make_account(owner_id=owner_a.id, company="List API Owner A Co")
    account_b = await make_account(owner_id=owner_b.id, company="List API Owner B Co")
    contact_a = await make_contact(account_id=account_a.id, first_name="Owned By A API")
    await make_contact(account_id=account_b.id, first_name="Owned By B API")
    headers = auth_headers(owner_a)

    response = await client.get(CONTACTS_URL, params={"owner_id": owner_a.id}, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert [item["id"] for item in body["items"]] == [contact_a.id]


async def test_list_contacts_route_no_auth_header_returns_401(client: AsyncClient):
    response = await client.get(CONTACTS_URL)

    assert response.status_code == 401


async def test_list_contacts_to_export_returns_valid_xlsx(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact
):
    owner = await make_user(email="rep-export-contact@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="List Export Contact Co")
    await make_contact(account_id=account.id, first_name="List", last_name="Export")
    headers = auth_headers(owner)

    response = await client.get(CONTACTS_URL, params={"to_export": "true"}, headers=headers)

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "contacts.xlsx" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == ["Name", "Email", "Phone", "Job Title", "Account", "Owner", "Tier", "Primary"]
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert any(row[0] == "List Export" for row in data_rows)


async def test_get_contact_to_export_returns_contact_and_deals_sheets(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact
):
    owner = await make_user(email="rep-export-single-contact@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Single Export Contact Co")
    contact = await make_contact(account_id=account.id, first_name="Single", last_name="Export")
    headers = auth_headers(owner)

    response = await client.get(
        f"{CONTACTS_URL}/{contact.id}", params={"to_export": "true"}, headers=headers
    )

    assert response.status_code == 200
    assert f"contact_{contact.id}.xlsx" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Contact", "Deals"]
    contact_sheet = workbook["Contact"]
    field_col = [cell.value for cell in contact_sheet["A"]]
    assert "Email" in field_col


# --- GET /contacts/{id}/overview -----------------------------------------------


async def test_get_contact_overview_route_returns_200(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact
):
    owner = await make_user(email="owner-overview-api@example.com", role=UserRole.SALES_REP, first_name="Karthick")
    headers = auth_headers(owner)
    account = await make_account(owner_id=owner.id, company="Nexbridge Tech", tier="gold")
    contact = await make_contact(account_id=account.id, first_name="Sarah", job_title="CTO", is_primary=True)

    response = await client.get(f"{CONTACTS_URL}/{contact.id}/overview", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == contact.id
    assert body["first_name"] == "Sarah"
    assert body["job_title"] == "CTO"
    assert body["is_primary"] is True
    assert body["account_id"] == account.id
    assert body["account_name"] == "Nexbridge Tech"
    assert body["owner_id"] == owner.id
    assert body["owner_name"] == "Karthick"
    assert body["tier"] == "gold"
    assert body["deal_count"] == 0
    assert body["task_count"] is None
    assert body["log_count"] is None
    assert body["tags"] is None
    assert body["about"] is None


async def test_get_contact_overview_route_returns_404_for_nonexistent_id(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-overview-404-api@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/999999/overview", headers=headers)

    assert response.status_code == 404


# --- GET /contacts/{id}/deals --------------------------------------------------


async def test_list_contact_deals_route_returns_200(
    client: AsyncClient, make_user, auth_headers, make_account, make_contact, make_deal
):
    owner = await make_user(email="owner-contact-deals-api@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(owner)
    account = await make_account(owner_id=owner.id, company="Contact Deals API Co")
    contact = await make_contact(account_id=account.id)
    deal = await make_deal(
        account_id=account.id, owner_id=owner.id, deal_name="Contact Deal API", contact_ids=[contact.id]
    )

    response = await client.get(f"{CONTACTS_URL}/{contact.id}/deals", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert [item["id"] for item in body] == [deal.id]


async def test_list_contact_deals_route_returns_404_for_nonexistent_id(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-contact-deals-404-api@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/999999/deals", headers=headers)

    assert response.status_code == 404


# --- GET /contacts/import/template -------------------------------------------


async def test_download_contact_import_template_xlsx(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-template-xlsx@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/import/template?format=xlsx", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"]


async def test_download_contact_import_template_csv(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-template-csv@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/import/template?format=csv", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]


async def test_download_contact_import_template_defaults_to_xlsx(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-template-default@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{CONTACTS_URL}/import/template", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")


# --- POST /contacts/import ----------------------------------------------------


async def test_upload_contact_import_csv_creates_contacts(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-upload-csv@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    csv_content = (
        b"first_name,last_name,email,phone,alternate_phone,job_title,linkedin_url\n"
        b"Jane,Doe,jane.upload.csv@acme.com,+1 555 0000,,VP,https://linkedin.com/in/jane\n"
    )

    response = await client.post(
        f"{CONTACTS_URL}/import",
        headers=headers,
        files={"file": ("contacts.csv", csv_content, "text/csv")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert body["errors"] == []
    contacts = await client.get(f"{CONTACTS_URL}?search=Jane", headers=headers)
    created = next(item for item in contacts.json()["items"] if item["email"] == "jane.upload.csv@acme.com")
    assert created["job_title"] == "VP"


async def test_upload_contact_import_xlsx_creates_contacts(client: AsyncClient, make_user, auth_headers):
    import io

    from openpyxl import Workbook

    rep = await make_user(email="rep-contact-upload-xlsx@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["first_name", "last_name", "email", "phone", "alternate_phone", "job_title", "linkedin_url"])
    sheet.append(["Jane", "Doe", "jane.upload.xlsx@acme.com", "", "", "VP", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = await client.post(
        f"{CONTACTS_URL}/import",
        headers=headers,
        files={
            "file": (
                "contacts.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert body["errors"] == []
    contacts = await client.get(f"{CONTACTS_URL}?search=Jane", headers=headers)
    created = next(item for item in contacts.json()["items"] if item["email"] == "jane.upload.xlsx@acme.com")
    assert created["job_title"] == "VP"


async def test_upload_contact_import_rejects_bad_extension(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-upload-bad-ext@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        f"{CONTACTS_URL}/import",
        headers=headers,
        files={"file": ("contacts.txt", b"not a real file", "text/plain")},
    )

    assert response.status_code == 400


async def test_upload_contact_import_rejects_corrupt_xlsx_content(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-contact-upload-corrupt@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        f"{CONTACTS_URL}/import",
        headers=headers,
        files={
            "file": (
                "contacts.xlsx",
                b"not actually an xlsx file",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 400
