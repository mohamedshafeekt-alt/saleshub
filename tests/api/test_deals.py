"""HTTP-level contract for /api/v1/deals.

Covers: 201 create as each allowed role, 422 missing required field
(deal_name/account_id/owner_id/stage_id), 401 no auth, 403 for Delivery SME
on create/list and for a non-owning Sales Rep on get/patch/delete/
stage-history, 404 for creating against a nonexistent account/stage and for
a missing deal id, 200 list scoped by role and by each filter (owner_id/
stage_id/search), sort_by/sort_dir, board view grouping, 200 partial PATCH,
400 PATCH that would leave a cold-stage deal with no cold_reason on record,
200 PATCH that sets a cold stage together with cold_reason in the same
request, 204 DELETE followed by a 404 GET, GET /deals/{id}/stage-history
reflecting stage transitions in order after PATCHes, xlsx export, and the
generic-patch route (happy path + 404 + invalid table + invalid field), plus
Deal Activities (create/list/update/delete under /deals/{id}/activities) and
Deal Documents (upload/list/delete under /deals/{id}/documents).
"""

import openpyxl
from httpx import AsyncClient

from app.models.enums import LeadTier
from tests.support.roles import UserRole

DEALS_URL = "/api/v1/deals"


def _deal_payload(**overrides) -> dict:
    payload = {
        "deal_name": "Acme Expansion",
        "account_id": overrides.pop("account_id"),
        "owner_id": overrides.pop("owner_id"),
        "stage_id": overrides.pop("stage_id"),
    }
    payload.update(overrides)
    return payload


async def test_create_deal_as_sales_rep_returns_201(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    rep = await make_user(email="rep-create-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Rep Deal Co")
    stage = await make_deal_stage(name="Received Requirement")
    headers = auth_headers(rep)

    response = await client.post(
        DEALS_URL,
        json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id, deal_name="Rep Deal"),
        headers=headers,
    )

    assert response.status_code == 201
    body = response.json()
    assert body["deal_name"] == "Rep Deal"
    assert body["account_id"] == account.id
    assert body["owner_id"] == rep.id
    assert body["currency"] == "USD"
    assert body["stage_id"] == stage.id
    assert "id" in body


async def test_create_deal_as_sales_manager_returns_201(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    manager = await make_user(email="manager-create-deal@example.com", role=UserRole.SALES_MANAGER)
    rep = await make_user(email="rep-owned-by-manager-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Manager Deal Co")
    stage = await make_deal_stage()
    headers = auth_headers(manager)

    response = await client.post(
        DEALS_URL, json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id), headers=headers
    )

    assert response.status_code == 201


async def test_create_deal_as_admin_returns_201(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    admin = await make_user(email="admin-create-deal@example.com", role=UserRole.ADMIN)
    rep = await make_user(email="rep-owned-by-admin-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Admin Deal Co")
    stage = await make_deal_stage()
    headers = auth_headers(admin)

    response = await client.post(
        DEALS_URL, json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id), headers=headers
    )

    assert response.status_code == 201


async def test_create_deal_missing_deal_name_returns_422(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    rep = await make_user(email="rep-missing-name-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Missing Name Deal Co")
    stage = await make_deal_stage()
    headers = auth_headers(rep)
    payload = _deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id)
    del payload["deal_name"]

    response = await client.post(DEALS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_deal_missing_account_id_returns_422(
    client: AsyncClient, make_user, auth_headers, make_deal_stage
):
    rep = await make_user(email="rep-missing-account-deal@example.com", role=UserRole.SALES_REP)
    stage = await make_deal_stage()
    headers = auth_headers(rep)
    payload = _deal_payload(account_id=1, owner_id=rep.id, stage_id=stage.id)
    del payload["account_id"]

    response = await client.post(DEALS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_deal_missing_owner_id_returns_422(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    rep = await make_user(email="rep-missing-owner-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Missing Owner Deal Co")
    stage = await make_deal_stage()
    headers = auth_headers(rep)
    payload = _deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id)
    del payload["owner_id"]

    response = await client.post(DEALS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_deal_missing_stage_id_returns_422(
    client: AsyncClient, make_user, auth_headers, make_account
):
    rep = await make_user(email="rep-missing-stage-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Missing Stage Deal Co")
    headers = auth_headers(rep)
    payload = {"deal_name": "No Stage Deal", "account_id": account.id, "owner_id": rep.id}

    response = await client.post(DEALS_URL, json=payload, headers=headers)

    assert response.status_code == 422


async def test_create_deal_no_auth_header_returns_401(
    client: AsyncClient, make_user, make_account, make_deal_stage
):
    rep = await make_user(email="rep-no-auth-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="No Auth Deal Co")
    stage = await make_deal_stage()

    response = await client.post(
        DEALS_URL, json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=stage.id)
    )

    assert response.status_code == 401


async def test_create_deal_as_delivery_sme_returns_403(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    sme = await make_user(email="sme-create-deal@example.com", role=UserRole.DELIVERY_SME)
    other_rep = await make_user(email="rep-for-sme-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=other_rep.id, company="SME Deal Co")
    stage = await make_deal_stage()
    headers = auth_headers(sme)

    response = await client.post(
        DEALS_URL,
        json=_deal_payload(account_id=account.id, owner_id=other_rep.id, stage_id=stage.id),
        headers=headers,
    )

    assert response.status_code == 403


async def test_create_deal_returns_404_for_nonexistent_account(
    client: AsyncClient, make_user, auth_headers, make_deal_stage
):
    rep = await make_user(email="rep-create-404-deal@example.com", role=UserRole.SALES_REP)
    stage = await make_deal_stage()
    headers = auth_headers(rep)

    response = await client.post(
        DEALS_URL, json=_deal_payload(account_id=999_999, owner_id=rep.id, stage_id=stage.id), headers=headers
    )

    assert response.status_code == 404


async def test_create_deal_returns_404_for_nonexistent_stage(
    client: AsyncClient, make_user, auth_headers, make_account
):
    rep = await make_user(email="rep-create-404-stage-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="404 Stage Deal Co")
    headers = auth_headers(rep)

    response = await client.post(
        DEALS_URL,
        json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=999_999),
        headers=headers,
    )

    assert response.status_code == 404


async def test_list_deals_as_delivery_sme_returns_403(client: AsyncClient, make_user, auth_headers):
    sme = await make_user(email="sme-list-deal@example.com", role=UserRole.DELIVERY_SME)
    headers = auth_headers(sme)

    response = await client.get(DEALS_URL, headers=headers)

    assert response.status_code == 403


async def test_list_deals_sales_rep_only_sees_own_deals(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep_a = await make_user(email="rep-list-a-deal@example.com", role=UserRole.SALES_REP)
    rep_b = await make_user(email="rep-list-b-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep_a.id, company="List Deal Co A")
    own_deal = await make_deal(account_id=account.id, owner_id=rep_a.id, deal_name="Own List Deal")
    await make_deal(account_id=account.id, owner_id=rep_b.id, deal_name="Other List Deal")
    headers = auth_headers(rep_a)

    response = await client.get(DEALS_URL, params={"owner_id": rep_b.id}, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert body["view"] == "list"
    assert [deal["id"] for deal in body["items"]] == [own_deal.id]


async def test_list_deals_manager_sees_all_deals(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    manager = await make_user(email="manager-list-deal@example.com", role=UserRole.SALES_MANAGER)
    rep_a = await make_user(email="rep-list-c-deal@example.com", role=UserRole.SALES_REP)
    rep_b = await make_user(email="rep-list-d-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep_a.id, company="List Deal Co B")
    deal_a = await make_deal(account_id=account.id, owner_id=rep_a.id, deal_name="List Deal A")
    deal_b = await make_deal(account_id=account.id, owner_id=rep_b.id, deal_name="List Deal B")
    headers = auth_headers(manager)

    response = await client.get(DEALS_URL, headers=headers)

    assert response.status_code == 200
    ids = {deal["id"] for deal in response.json()["items"]}
    assert {deal_a.id, deal_b.id} <= ids


async def test_list_deals_total_reflects_full_filtered_count_not_page_size(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep = await make_user(email="rep-total-count-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Total Count Deal Co")
    for i in range(3):
        await make_deal(account_id=account.id, owner_id=rep.id, deal_name=f"Total Count Deal {i}")
    headers = auth_headers(rep)

    response = await client.get(DEALS_URL, params={"limit": 2, "offset": 0}, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert len(body["items"]) == 2
    assert body["total"] == 3


async def test_list_deals_filters_by_stage_id(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal, make_deal_stage
):
    rep = await make_user(email="rep-filter-stage-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Filter Stage Deal Co")
    stage_a = await make_deal_stage(name="Eval Filter Stage")
    stage_b = await make_deal_stage(name="Proposal Filter Stage")
    eval_deal = await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Eval Filter Deal", stage_id=stage_a.id
    )
    await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Proposal Filter Deal", stage_id=stage_b.id
    )
    headers = auth_headers(rep)

    response = await client.get(DEALS_URL, params={"stage_id": stage_a.id}, headers=headers)

    assert response.status_code == 200
    assert [deal["id"] for deal in response.json()["items"]] == [eval_deal.id]


async def test_list_deals_filters_by_tier(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep = await make_user(email="rep-filter-tier-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Filter Tier Deal Co")
    gold_deal = await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Gold Tier Deal", tier=LeadTier.GOLD
    )
    await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Silver Tier Deal", tier=LeadTier.SILVER
    )
    headers = auth_headers(rep)

    response = await client.get(DEALS_URL, params={"tier": "gold"}, headers=headers)

    assert response.status_code == 200
    assert [deal["id"] for deal in response.json()["items"]] == [gold_deal.id]


async def test_list_deals_filters_by_search_matches_deal_name_or_account(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep = await make_user(email="rep-filter-search-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Filter Search Deal Co")
    match = await make_deal(account_id=account.id, owner_id=rep.id, deal_name="Searchable Deal Inc")
    await make_deal(account_id=account.id, owner_id=rep.id, deal_name="Nothing Deal")
    account_match = await make_account(owner_id=rep.id, company="Zephyr Industries")
    account_match_deal = await make_deal(
        account_id=account_match.id, owner_id=rep.id, deal_name="Some Other Deal"
    )
    headers = auth_headers(rep)

    response = await client.get(DEALS_URL, params={"search": "searchable"}, headers=headers)
    assert response.status_code == 200
    assert [deal["id"] for deal in response.json()["items"]] == [match.id]

    response = await client.get(DEALS_URL, params={"search": "zephyr"}, headers=headers)
    assert response.status_code == 200
    assert [deal["id"] for deal in response.json()["items"]] == [account_match_deal.id]


async def test_list_deals_sort_by_value_asc(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep = await make_user(email="rep-sort-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Sort Deal Co")
    low = await make_deal(account_id=account.id, owner_id=rep.id, deal_name="Low Deal", value=10.0)
    high = await make_deal(account_id=account.id, owner_id=rep.id, deal_name="High Deal", value=99.0)
    headers = auth_headers(rep)

    response = await client.get(
        DEALS_URL, params={"sort_by": "value", "sort_dir": "asc"}, headers=headers
    )

    assert response.status_code == 200
    assert [deal["id"] for deal in response.json()["items"]] == [low.id, high.id]


async def test_list_deals_board_view_groups_by_stage(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal, make_deal_stage
):
    rep = await make_user(email="rep-board-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Board View Deal Co")
    stage_a = await make_deal_stage(name="Board View Stage A", sort_order=0)
    stage_b = await make_deal_stage(name="Board View Stage B", sort_order=1)
    await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Board Deal A1", stage_id=stage_a.id, value=100.0
    )
    await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Board Deal A2", stage_id=stage_a.id, value=50.0
    )
    await make_deal(
        account_id=account.id, owner_id=rep.id, deal_name="Board Deal B1", stage_id=stage_b.id, value=25.0
    )
    headers = auth_headers(rep)

    response = await client.get(DEALS_URL, params={"view": "board"}, headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert body["view"] == "board"
    assert body["items"] is None
    columns = body["columns"]
    assert [col["stage_id"] for col in columns] == [stage_a.id, stage_b.id]
    assert columns[0]["stage_name"] == "Board View Stage A"
    assert columns[0]["total_value"] == 150.0
    assert len(columns[0]["deals"]) == 2
    assert columns[1]["total_value"] == 25.0


async def test_get_deal_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-get-404-deal@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{DEALS_URL}/999999", headers=headers)

    assert response.status_code == 404


async def test_get_deal_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-owns-get-deal@example.com", role=UserRole.SALES_REP)
    other_rep = await make_user(email="rep-not-owner-get-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Get Forbidden Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Get Forbidden Deal")
    headers = auth_headers(other_rep)

    response = await client.get(f"{DEALS_URL}/{deal.id}", headers=headers)

    assert response.status_code == 403


async def test_get_deal_returns_200_for_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-owns-get-ok-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Get Ok Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Get Ok Deal")
    headers = auth_headers(owner)

    response = await client.get(f"{DEALS_URL}/{deal.id}", headers=headers)

    assert response.status_code == 200
    assert response.json()["id"] == deal.id


async def test_update_deal_partial_patch_returns_200(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-patch-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Patch Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Old Patch Deal", value=1000.0)
    headers = auth_headers(owner)

    response = await client.patch(
        f"{DEALS_URL}/{deal.id}", json={"deal_name": "New Patch Deal"}, headers=headers
    )

    assert response.status_code == 200
    body = response.json()
    assert body["deal_name"] == "New Patch Deal"
    assert body["value"] == 1000.0


async def test_update_deal_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-patch-404-deal@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.patch(f"{DEALS_URL}/999999", json={"deal_name": "New"}, headers=headers)

    assert response.status_code == 404


async def test_update_deal_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-owns-patch-deal@example.com", role=UserRole.SALES_REP)
    other_rep = await make_user(email="rep-not-owner-patch-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Patch Forbidden Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Patch Forbidden Deal")
    headers = auth_headers(other_rep)

    response = await client.patch(f"{DEALS_URL}/{deal.id}", json={"deal_name": "New"}, headers=headers)

    assert response.status_code == 403


async def test_update_deal_returns_400_when_stage_cold_without_reason(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal, make_deal_stage
):
    owner = await make_user(email="rep-patch-cold-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Patch Cold Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Patch Cold Deal")
    cold_stage = await make_deal_stage(name="Patch Cold Stage", is_cold=True)
    headers = auth_headers(owner)

    response = await client.patch(
        f"{DEALS_URL}/{deal.id}", json={"stage_id": cold_stage.id}, headers=headers
    )

    assert response.status_code == 400


async def test_update_deal_returns_200_when_cold_reason_provided_with_stage(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal, make_deal_stage
):
    owner = await make_user(email="rep-patch-cold-ok-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Patch Cold Ok Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Patch Cold Ok Deal")
    cold_stage = await make_deal_stage(name="Patch Cold Ok Stage", is_cold=True)
    headers = auth_headers(owner)

    response = await client.patch(
        f"{DEALS_URL}/{deal.id}",
        json={"stage_id": cold_stage.id, "cold_reason": "Lost budget"},
        headers=headers,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["stage_id"] == cold_stage.id
    assert body["cold_reason"] == "Lost budget"


async def test_create_deal_returns_400_when_stage_cold_without_reason(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal_stage
):
    rep = await make_user(email="rep-create-cold-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep.id, company="Create Cold Deal Co")
    cold_stage = await make_deal_stage(name="Create Cold Stage", is_cold=True)
    headers = auth_headers(rep)

    response = await client.post(
        DEALS_URL,
        json=_deal_payload(account_id=account.id, owner_id=rep.id, stage_id=cold_stage.id),
        headers=headers,
    )

    assert response.status_code == 400


async def test_delete_deal_returns_204(client: AsyncClient, make_user, auth_headers, make_account, make_deal):
    owner = await make_user(email="rep-delete-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Delete Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Delete Me Deal")
    headers = auth_headers(owner)

    response = await client.delete(f"{DEALS_URL}/{deal.id}", headers=headers)

    assert response.status_code == 204

    follow_up = await client.get(f"{DEALS_URL}/{deal.id}", headers=headers)
    assert follow_up.status_code == 404


async def test_delete_deal_returns_404_for_nonexistent_id(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-delete-404-deal@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.delete(f"{DEALS_URL}/999999", headers=headers)

    assert response.status_code == 404


async def test_delete_deal_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-owns-delete-deal@example.com", role=UserRole.SALES_REP)
    other_rep = await make_user(email="rep-not-owner-delete-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Delete Forbidden Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Delete Forbidden Deal")
    headers = auth_headers(other_rep)

    response = await client.delete(f"{DEALS_URL}/{deal.id}", headers=headers)

    assert response.status_code == 403


# --- GET /deals/{deal_id}/stage-history --------------------------------------


async def test_stage_history_returns_200_reflecting_transitions_after_patches(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal, make_deal_stage
):
    owner = await make_user(email="rep-history-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="History Deal Co")
    stage_1 = await make_deal_stage(name="History Stage 1")
    stage_2 = await make_deal_stage(name="History Stage 2")
    stage_3 = await make_deal_stage(name="History Stage 3")
    deal = await make_deal(
        account_id=account.id, owner_id=owner.id, deal_name="History Deal", stage_id=stage_1.id
    )
    headers = auth_headers(owner)

    await client.patch(f"{DEALS_URL}/{deal.id}", json={"stage_id": stage_2.id}, headers=headers)
    await client.patch(f"{DEALS_URL}/{deal.id}", json={"stage_id": stage_3.id}, headers=headers)

    response = await client.get(f"{DEALS_URL}/{deal.id}/stage-history", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert [row["to_stage_id"] for row in body] == [stage_2.id, stage_3.id]


async def test_stage_history_returns_404_for_nonexistent_deal(client: AsyncClient, make_user, auth_headers):
    rep = await make_user(email="rep-history-404-deal@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.get(f"{DEALS_URL}/999999/stage-history", headers=headers)

    assert response.status_code == 404


async def test_stage_history_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-owns-history-deal@example.com", role=UserRole.SALES_REP)
    other_rep = await make_user(email="rep-not-owner-history-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="History Forbidden Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="History Forbidden Deal")
    headers = auth_headers(other_rep)

    response = await client.get(f"{DEALS_URL}/{deal.id}/stage-history", headers=headers)

    assert response.status_code == 403


# --- GET /deals?to_export=true --------------------------------------------


async def test_list_deals_to_export_returns_valid_xlsx_with_expected_rows(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-export-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Export Xlsx Deal Co")
    await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Export Xlsx Deal", value=42.0)
    headers = auth_headers(owner)

    response = await client.get(DEALS_URL, params={"to_export": "true"}, headers=headers)

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "deals.xlsx" in response.headers["content-disposition"]

    import io

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == [
        "Deal Name",
        "Account",
        "Contact",
        "Value",
        "Currency",
        "Stage",
        "Tier",
        "Owner",
        "Expected Close Date",
        "Cold Reason",
    ]
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert any(row[0] == "Export Xlsx Deal" for row in data_rows)


async def test_deals_export_route_no_longer_exists(client: AsyncClient, make_user, auth_headers):
    """"/export" now falls through to GET /deals/{deal_id} and fails int
    path-param validation (422), since the dedicated /export route is gone."""
    owner = await make_user(email="rep-old-export-deal-route@example.com", role=UserRole.SALES_REP)
    response = await client.get(f"{DEALS_URL}/export", headers=auth_headers(owner))
    assert response.status_code == 422


async def test_list_deals_to_export_scopes_to_requester_for_non_view_all_role(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    rep_a = await make_user(email="rep-a-export-xlsx@example.com", role=UserRole.SALES_REP)
    rep_b = await make_user(email="rep-b-export-xlsx@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=rep_a.id, company="Export Scope Xlsx Co")
    await make_deal(account_id=account.id, owner_id=rep_a.id, deal_name="Own Export Xlsx Deal")
    await make_deal(account_id=account.id, owner_id=rep_b.id, deal_name="Other Export Xlsx Deal")
    headers = auth_headers(rep_a)

    response = await client.get(DEALS_URL, params={"to_export": "true"}, headers=headers)

    import io

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    names = {row[0] for row in data_rows}
    assert names == {"Own Export Xlsx Deal"}


async def test_get_deal_to_export_returns_deal_and_stage_history_sheets(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-export-single-deal@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Single Export Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Single Export Deal")
    headers = auth_headers(owner)

    response = await client.get(f"{DEALS_URL}/{deal.id}", params={"to_export": "true"}, headers=headers)

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"deal_{deal.id}.xlsx" in response.headers["content-disposition"]

    import io

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Deal", "Stage History"]
    deal_sheet = workbook["Deal"]
    field_col = [cell.value for cell in deal_sheet["A"]]
    assert "Deal Name" in field_col


# --- PATCH /deals/generic-patch ------------------------------------------


async def test_generic_patch_updates_allowed_field(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-generic-patch@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Generic Patch Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Generic Patch Deal")
    headers = auth_headers(owner)

    response = await client.patch(
        f"{DEALS_URL}/generic-patch",
        json={"table": "deals", "record_id": deal.id, "field": "deal_name", "value": "Patched Name"},
        headers=headers,
    )

    assert response.status_code == 200
    body = response.json()
    assert body == {"id": deal.id, "field": "deal_name", "value": "Patched Name"}

    follow_up = await client.get(f"{DEALS_URL}/{deal.id}", headers=headers)
    assert follow_up.json()["deal_name"] == "Patched Name"


async def test_generic_patch_returns_404_for_missing_record(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-generic-patch-404@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.patch(
        f"{DEALS_URL}/generic-patch",
        json={"table": "deals", "record_id": 999_999, "field": "deal_name", "value": "X"},
        headers=headers,
    )

    assert response.status_code == 404


async def test_generic_patch_returns_400_for_disallowed_table(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-generic-patch-table@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.patch(
        f"{DEALS_URL}/generic-patch",
        json={"table": "users", "record_id": 1, "field": "email", "value": "x@example.com"},
        headers=headers,
    )

    assert response.status_code == 400


async def test_generic_patch_returns_400_for_disallowed_field(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner = await make_user(email="rep-generic-patch-field@example.com", role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company="Generic Patch Field Deal Co")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name="Generic Patch Field Deal")
    headers = auth_headers(owner)

    response = await client.patch(
        f"{DEALS_URL}/generic-patch",
        json={"table": "deals", "record_id": deal.id, "field": "not_a_real_column", "value": "x"},
        headers=headers,
    )

    assert response.status_code == 400


async def _make_owned_deal(make_user, make_account, make_deal, *, email: str):
    owner = await make_user(email=email, role=UserRole.SALES_REP)
    account = await make_account(owner_id=owner.id, company=f"Acme {email}")
    deal = await make_deal(account_id=account.id, owner_id=owner.id, deal_name=f"Deal {email}")
    return owner, deal


async def test_create_deal_activity_returns_201(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-deal-activity-create@example.com"
    )
    headers = auth_headers(owner)

    response = await client.post(
        f"{DEALS_URL}/{deal.id}/activities",
        json={"type": "call", "note": "Discussed pricing"},
        headers=headers,
    )

    assert response.status_code == 201
    body = response.json()
    assert body["deal_id"] == deal.id
    assert body["type"] == "call"
    assert body["note"] == "Discussed pricing"
    assert body["created_by"] == owner.id


async def test_create_deal_activity_returns_404_for_nonexistent_deal(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-deal-activity-404@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        f"{DEALS_URL}/999999/activities", json={"type": "note", "note": "x"}, headers=headers
    )

    assert response.status_code == 404


async def test_create_deal_activity_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-owns-deal-activity@example.com"
    )
    other_rep = await make_user(email="rep-not-owner-deal-activity@example.com", role=UserRole.SALES_REP)

    response = await client.post(
        f"{DEALS_URL}/{deal.id}/activities",
        json={"type": "note", "note": "x"},
        headers=auth_headers(other_rep),
    )

    assert response.status_code == 403


async def test_list_deal_activities_returns_200_filtered_by_type(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-list-deal-activities@example.com"
    )
    headers = auth_headers(owner)
    await client.post(
        f"{DEALS_URL}/{deal.id}/activities", json={"type": "call", "note": "call note"}, headers=headers
    )
    await client.post(
        f"{DEALS_URL}/{deal.id}/activities", json={"type": "note", "note": "note note"}, headers=headers
    )

    response = await client.get(f"{DEALS_URL}/{deal.id}/activities?types=call", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert len(body) == 1
    assert body[0]["type"] == "call"
    assert body[0]["created_by_name"]


async def test_update_deal_activity_returns_200(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-update-deal-activity@example.com"
    )
    headers = auth_headers(owner)
    created = (
        await client.post(
            f"{DEALS_URL}/{deal.id}/activities", json={"type": "note", "note": "original"}, headers=headers
        )
    ).json()

    response = await client.patch(
        f"{DEALS_URL}/{deal.id}/activities/{created['id']}",
        json={"note": "revised"},
        headers=headers,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["note"] == "revised"
    assert body["updated_by"] == owner.id
    assert body["updated_by_name"]


async def test_delete_deal_activity_returns_204_then_404_on_refetch(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-delete-deal-activity@example.com"
    )
    admin = await make_user(email="admin-delete-deal-activity@example.com", role=UserRole.ADMIN)
    owner_headers = auth_headers(owner)
    created = (
        await client.post(
            f"{DEALS_URL}/{deal.id}/activities",
            json={"type": "note", "note": "to delete"},
            headers=owner_headers,
        )
    ).json()

    delete_response = await client.delete(
        f"{DEALS_URL}/{deal.id}/activities/{created['id']}", headers=auth_headers(admin)
    )

    assert delete_response.status_code == 204
    refetch = await client.patch(
        f"{DEALS_URL}/{deal.id}/activities/{created['id']}", json={"note": "x"}, headers=owner_headers
    )
    assert refetch.status_code == 404


async def test_delete_deal_activity_returns_403_for_deal_owner(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-delete-deal-activity-forbidden@example.com"
    )
    headers = auth_headers(owner)
    created = (
        await client.post(
            f"{DEALS_URL}/{deal.id}/activities",
            json={"type": "note", "note": "cannot delete"},
            headers=headers,
        )
    ).json()

    response = await client.delete(f"{DEALS_URL}/{deal.id}/activities/{created['id']}", headers=headers)

    assert response.status_code == 403


async def test_upload_deal_document_returns_201(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-upload-deal-doc@example.com"
    )
    headers = auth_headers(owner)

    response = await client.post(
        f"{DEALS_URL}/{deal.id}/documents",
        files={"file": ("proposal.pdf", b"%PDF-1.4 fake", "application/pdf")},
        headers=headers,
    )

    assert response.status_code == 201
    body = response.json()
    assert body["deal_id"] == deal.id
    assert body["file_name"] == "proposal.pdf"
    assert body["content_type"] == "application/pdf"
    assert body["uploaded_by"] == owner.id
    assert body["file_url"].startswith("/media/deal_documents/")

    import pathlib

    pathlib.Path(body["file_url"].lstrip("/")).unlink()


async def test_upload_deal_document_rejects_unsupported_type_returns_400(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-upload-deal-doc-bad-type@example.com"
    )
    headers = auth_headers(owner)

    response = await client.post(
        f"{DEALS_URL}/{deal.id}/documents",
        files={"file": ("virus.exe", b"whatever", "application/x-msdownload")},
        headers=headers,
    )

    assert response.status_code == 400


async def test_upload_deal_document_returns_404_for_nonexistent_deal(
    client: AsyncClient, make_user, auth_headers
):
    rep = await make_user(email="rep-upload-deal-doc-404@example.com", role=UserRole.SALES_REP)
    headers = auth_headers(rep)

    response = await client.post(
        f"{DEALS_URL}/999999/documents",
        files={"file": ("x.pdf", b"x", "application/pdf")},
        headers=headers,
    )

    assert response.status_code == 404


async def test_upload_deal_document_returns_403_for_non_owning_sales_rep(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-owns-deal-doc@example.com"
    )
    other_rep = await make_user(email="rep-not-owner-deal-doc@example.com", role=UserRole.SALES_REP)

    response = await client.post(
        f"{DEALS_URL}/{deal.id}/documents",
        files={"file": ("x.pdf", b"x", "application/pdf")},
        headers=auth_headers(other_rep),
    )

    assert response.status_code == 403


async def test_list_and_delete_deal_documents(
    client: AsyncClient, make_user, auth_headers, make_account, make_deal
):
    owner, deal = await _make_owned_deal(
        make_user, make_account, make_deal, email="rep-list-delete-deal-doc@example.com"
    )
    headers = auth_headers(owner)
    created = (
        await client.post(
            f"{DEALS_URL}/{deal.id}/documents",
            files={"file": ("proposal.pdf", b"content", "application/pdf")},
            headers=headers,
        )
    ).json()

    list_response = await client.get(f"{DEALS_URL}/{deal.id}/documents", headers=headers)
    assert list_response.status_code == 200
    assert [doc["id"] for doc in list_response.json()] == [created["id"]]

    delete_response = await client.delete(
        f"{DEALS_URL}/{deal.id}/documents/{created['id']}", headers=headers
    )
    assert delete_response.status_code == 204

    list_after_delete = await client.get(f"{DEALS_URL}/{deal.id}/documents", headers=headers)
    assert list_after_delete.json() == []
